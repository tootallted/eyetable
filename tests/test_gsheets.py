import datetime as dt
import io

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from eyetable import gsheets
from eyetable.api import create_app
from eyetable.engine import Engine

TOKEN = "test-token-0123456789"
AUTH = {"Authorization": f"Bearer {TOKEN}"}
SHEET_ID = "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789"
LINK = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit#gid=0"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(tabs: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in tabs.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


LEDGER = {
    "Q1 Actuals": [
        ["Account", "Amount", "Posted", "Code", None],
        ["Rent", 1200.5, dt.datetime(2026, 1, 31), "007", None],
        ["Payroll", 5400, dt.datetime(2026, 2, 28), "010"],
    ],
    "Q2 Actuals": [["Account", "Amount"], ["Rent", 1250]],
    "Notes": [],
}


def fake_fetch(tabs=LEDGER, title="General Ledger 2026"):
    def fetch(url):
        sid, published = gsheets.parse_sheet_url(url)
        return gsheets.Workbook(title, sid, published, gsheets.sheet_link(sid, published),
                                gsheets.parse_workbook(make_xlsx(tabs)))
    return fetch


def test_parse_links():
    assert gsheets.parse_sheet_url(LINK) == (SHEET_ID, False)
    assert gsheets.parse_sheet_url(f"https://docs.google.com/spreadsheets/u/1/d/{SHEET_ID}/view") == (SHEET_ID, False)
    assert gsheets.parse_sheet_url(f"https://docs.google.com/spreadsheets/d/e/2PACX-{SHEET_ID}/pubhtml") == (f"2PACX-{SHEET_ID}", True)
    assert gsheets.parse_sheet_url(SHEET_ID) == (SHEET_ID, False)
    for bad in ("https://example.com/spreadsheets/d/x", "hello", "https://docs.google.com/document/d/" + SHEET_ID):
        with pytest.raises(gsheets.SheetError):
            gsheets.parse_sheet_url(bad)


def test_title_from_download_header():
    header = "attachment; filename=\"GeneralLedger.xlsx\"; filename*=UTF-8''General%20Ledger%202026.xlsx"
    assert gsheets.title_from_disposition(header) == "General Ledger 2026"
    assert gsheets.title_from_disposition('attachment; filename="Budget.xlsx"') == "Budget"
    assert gsheets.title_from_disposition(None) is None


def test_names():
    assert gsheets.source_name("General Ledger 2026", "Q1 Actuals") == "general_ledger_2026_q1_actuals"
    assert gsheets.source_name("2026 Budget", "Main") == "s_2026_budget_main"
    assert len(gsheets.source_name("x" * 80, "y" * 80)) == 63


def test_tabs_are_parsed_and_typed():
    tabs = gsheets.parse_workbook(make_xlsx(LEDGER))
    q1 = tabs[0]
    assert (q1.title, q1.headers) == ("Q1 Actuals", ["Account", "Amount", "Posted", "Code"])
    assert q1.records[0] == {"Account": "Rent", "Amount": 1200.5, "Posted": dt.date(2026, 1, 31), "Code": "007"}
    assert tabs[2].headers == [] and tabs[2].records == []


def test_fetch_through_google(monkeypatch):
    xlsx = make_xlsx(LEDGER)
    seen = []

    def handler(request):
        seen.append(str(request.url))
        return httpx.Response(200, content=xlsx, headers={
            "content-type": XLSX_TYPE,
            "content-disposition": "attachment; filename*=UTF-8''General%20Ledger%202026.xlsx",
        })

    book = gsheets.fetch_workbook(LINK, client=httpx.Client(transport=httpx.MockTransport(handler)))
    assert book.title == "General Ledger 2026" and [t.title for t in book.tabs] == ["Q1 Actuals", "Q2 Actuals", "Notes"]
    assert seen == [f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"]


def test_private_sheet_explains_sharing():
    def sign_in_page(request):
        return httpx.Response(200, text="<html>Sign in</html>", headers={"content-type": "text/html"})

    with pytest.raises(gsheets.SheetError, match="Anyone with the link"):
        gsheets.fetch_workbook(LINK, client=httpx.Client(transport=httpx.MockTransport(sign_in_page)))


def test_every_tab_becomes_a_source(store):
    engine = Engine(store)
    load = engine.ingest_gsheet(LINK, fetch=fake_fetch())
    assert [(r.dataset_key, r.status, r.rows) for r in load.results] == [
        ("Q1 Actuals", "succeeded", 2), ("Q2 Actuals", "succeeded", 1), ("Notes", "skipped", 0),
    ]
    q1 = store.get_source("general_ledger_2026_q1_actuals")
    assert q1.type == "gsheet" and q1.config["spreadsheet_id"] == SHEET_ID and q1.config["tab_title"] == "Q1 Actuals"
    rows = store.read_table("general_ledger_2026_q1_actuals__data")
    assert [r["ET_id"] for r in rows] == [1, 2]
    assert rows[0]["ET_source"] == "general_ledger_2026_q1_actuals - General Ledger 2026 / Q1 Actuals"
    assert rows[0]["code"] == "007" and rows[0]["posted"] == dt.date(2026, 1, 31)
    assert not store.source_exists("general_ledger_2026_notes")


def test_reloading_replaces_or_appends_and_picks_up_new_tabs(store):
    engine = Engine(store)
    engine.ingest_gsheet(LINK, fetch=fake_fetch())
    more = {**LEDGER, "Q2 Actuals": [["Account", "Amount"], ["Rent", 1250], ["Travel", 300]], "Q3 Actuals": [["Account"], ["Rent"]]}

    plan = engine.plan_gsheet(fake_fetch(more)(LINK))
    assert [(t.source, t.exists) for t in plan.tabs if not t.problem] == [
        ("general_ledger_2026_q1_actuals", True), ("general_ledger_2026_q2_actuals", True), ("general_ledger_2026_q3_actuals", False),
    ]

    engine.ingest_gsheet(LINK, mode="append", fetch=fake_fetch(more))
    assert [r["ET_id"] for r in store.read_table("general_ledger_2026_q2_actuals__data")] == [1, 2, 3]
    assert len(store.read_table("general_ledger_2026_q3_actuals__data")) == 1  # new tab: created fresh

    engine.ingest_gsheet(LINK, mode="replace", fetch=fake_fetch(more))
    assert [r["ET_id"] for r in store.read_table("general_ledger_2026_q2_actuals__data")] == [1, 2]
    assert len(store.list_uploads("general_ledger_2026_q2_actuals")) == 3


def test_name_taken_by_another_source_is_skipped(store):
    engine = Engine(store)
    engine.ingest_upload("general_ledger_2026_q1_actuals", "x.csv", b"a\n1\n", create=True)
    load = engine.ingest_gsheet(LINK, fetch=fake_fetch())
    assert load.results[0].status == "skipped" and "already named" in load.results[0].error
    assert store.get_source("general_ledger_2026_q1_actuals").type == "file"
    assert load.results[1].status == "succeeded"


def test_api_inspect_load_and_query(store):
    client = TestClient(create_app(Engine(store), TOKEN, sheet_fetch=fake_fetch()))
    plan = client.post("/api/gsheets/inspect", headers=AUTH, json={"url": LINK}).json()
    assert plan["title"] == "General Ledger 2026"
    assert [(t["source"], t["exists"], t["problem"]) for t in plan["tabs"]] == [
        ("general_ledger_2026_q1_actuals", False, None),
        ("general_ledger_2026_q2_actuals", False, None),
        ("general_ledger_2026_notes", False, "the tab is empty"),
    ]
    assert not store.list_sources()  # inspecting changes nothing

    load = client.post("/api/gsheets", headers=AUTH, json={"url": LINK}).json()
    assert [r["status"] for r in load["results"]] == ["succeeded", "succeeded", "skipped"]

    summary = {s["name"]: s for s in client.get("/api/sources", headers=AUTH).json()}
    assert summary["general_ledger_2026_q1_actuals"]["type"] == "gsheet"
    assert summary["general_ledger_2026_q1_actuals"]["rows"] == 2

    preview = client.get("/api/sources/general_ledger_2026_q1_actuals/preview", headers=AUTH).json()
    assert preview["rows"][0][2] == "general_ledger_2026_q1_actuals - General Ledger 2026 / Q1 Actuals"

    q = client.post("/api/query", headers=AUTH, json={"sql": "SELECT account, amount FROM general_ledger_2026_q1_actuals ORDER BY amount"})
    assert q.json()["rows"] == [["Rent", 1200.5], ["Payroll", 5400.0]]

    assert client.post("/api/gsheets", headers=AUTH, json={"url": "https://example.com/x"}).status_code == 400
    assert client.post("/api/gsheets", json={"url": LINK}).status_code == 401
