"""Google Sheets: one eyetable source per tab of a spreadsheet.

Two ways in:

  * Signed out: the sheet is downloaded the way Google's own
    "Download > Microsoft Excel" works (an .xlsx export with every tab). No
    Google account is involved, so the sheet must be readable by link
    ("Anyone with the link can view", or published to the web).
  * Signed in: the page gets a short-lived, read-only Google access token for
    the person using it (Sign in with Google) and passes it along; the sheet
    is read through the Sheets API as that person, so private sheets they can
    open work too. The token is used for that one request and never stored.

Each tab becomes a source named "<file name>_<tab name>", e.g. the tab
"Q1 Actuals" in "General Ledger 2026" -> general_ledger_2026_q1_actuals.
"""

import datetime as dt
import io
import json
import re
import urllib.parse
from dataclasses import dataclass, field
from typing import Any

import httpx

from eyetable.connectors.file import unique_headers
from eyetable.identifiers import sanitize
from eyetable.identifiers import source_name as make_source_name

_SHEET_URL = re.compile(r"docs\.google\.com/spreadsheets/(?:u/\d+/)?d/(e/)?([A-Za-z0-9_-]{20,})")
_BARE_ID = re.compile(r"[A-Za-z0-9_-]{25,}")
_XLSX = "spreadsheetml"


SHEETS_API = "https://sheets.googleapis.com/v4/spreadsheets"
# Just what eyetable needs: titles, tab properties, and each cell's value and number format.
SHEETS_FIELDS = (
    "properties.title,"
    "sheets(properties(title,index,hidden),"
    "data(rowData(values(effectiveValue,effectiveFormat.numberFormat.type))))"
)
_SHEETS_EPOCH = dt.datetime(1899, 12, 30)  # day 0 of Sheets' date serial numbers
_ERROR_TEXT = {
    "DIVIDE_BY_ZERO": "#DIV/0!", "N_A": "#N/A", "VALUE": "#VALUE!", "REF": "#REF!",
    "NAME": "#NAME?", "NUM": "#NUM!", "NULL_VALUE": "#NULL!", "LOADING": "#LOADING",
}


class SheetError(ValueError):
    def __init__(self, message: str, needs_sign_in: bool = False):
        super().__init__(message)
        self.needs_sign_in = needs_sign_in  # signing in with Google may get past this


@dataclass
class Tab:
    title: str
    index: int
    headers: list[str]
    records: list[dict[str, Any]]
    hidden: bool = False


@dataclass
class Workbook:
    title: str
    spreadsheet_id: str
    published: bool
    url: str
    tabs: list[Tab] = field(default_factory=list)


def parse_sheet_url(url: str) -> tuple[str, bool]:
    """(spreadsheet id, is a published-to-web id) from a Sheets link or a bare id."""
    url = url.strip()
    match = _SHEET_URL.search(url)
    if match:
        return match.group(2), bool(match.group(1))
    if _BARE_ID.fullmatch(url):
        return url, False
    raise SheetError("that doesn't look like a Google Sheets link (https://docs.google.com/spreadsheets/d/...)")


def export_url(spreadsheet_id: str, published: bool) -> str:
    # Built from the id alone, never from the caller's URL, so only Google is ever fetched.
    if published:
        return f"https://docs.google.com/spreadsheets/d/e/{spreadsheet_id}/pub?output=xlsx"
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=xlsx"


def sheet_link(spreadsheet_id: str, published: bool) -> str:
    if published:
        return f"https://docs.google.com/spreadsheets/d/e/{spreadsheet_id}/pubhtml"
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit"


def title_from_disposition(header: str | None) -> str | None:
    """The file name Google puts on the download, without its extension."""
    if not header:
        return None
    star = re.search(r"filename\*\s*=\s*(?:UTF-8|utf-8)''([^;]+)", header)
    if star:
        name = urllib.parse.unquote(star.group(1).strip().strip('"'))
    else:
        plain = re.search(r'filename\s*=\s*"([^"]*)"', header) or re.search(r"filename\s*=\s*([^;]+)", header)
        if not plain:
            return None
        name = plain.group(1).strip()
    name = re.sub(r"\.xlsx$", "", name, flags=re.I).strip()
    return name or None


def source_name(file_title: str, tab_title: str) -> str:
    return make_source_name(f"{sanitize(file_title)}_{sanitize(tab_title)}")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_dates(records: list[dict[str, Any]], headers: list[str]) -> None:
    """Sheets dates arrive as midnight datetimes; a column of only those is a date column."""
    for name in headers:
        values = [r[name] for r in records if r[name] is not None]
        if values and all(isinstance(v, dt.datetime) and v.time() == dt.time() and v.tzinfo is None for v in values):
            for r in records:
                if r[name] is not None:
                    r[name] = r[name].date()


def tab_from_rows(title: str, index: int, hidden: bool, rows: list[list[Any]]) -> Tab:
    """A tab from its cell values: the first non-empty row is the header."""
    rows = [list(r) for r in rows if any(v is not None and _cell_text(v) != "" for v in r)]
    if not rows:
        return Tab(title, index, [], [], hidden)
    width = max(max((i + 1 for i, v in enumerate(r) if v is not None and _cell_text(v) != ""), default=0) for r in rows)
    headers = unique_headers([_cell_text(v) for v in rows[0]], width)
    records = []
    for r in rows[1:]:
        record = {}
        for i, name in enumerate(headers):
            v = r[i] if i < len(r) else None
            if isinstance(v, str):
                v = v if v.strip() else None
            record[name] = v
        records.append(record)
    _normalize_dates(records, headers)
    return Tab(title, index, headers, records, hidden)


def parse_workbook(content: bytes) -> list[Tab]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise SheetError(f"Google's export could not be read as a spreadsheet ({type(exc).__name__})") from None
    try:
        return [
            tab_from_rows(ws.title, index, ws.sheet_state != "visible", ws.iter_rows(values_only=True))
            for index, ws in enumerate(wb.worksheets)
        ]
    finally:
        wb.close()


def cell_value(cell: dict[str, Any]) -> Any:
    """A Sheets API cell (effectiveValue + number format type) as a Python value."""
    value = cell.get("effectiveValue")
    if not value:
        return None
    if "stringValue" in value:
        return value["stringValue"]
    if "boolValue" in value:
        return value["boolValue"]
    if "errorValue" in value:
        return _ERROR_TEXT.get(value["errorValue"].get("type"), "#ERROR!")
    number = value.get("numberValue")
    if number is None:
        return None
    kind = cell.get("effectiveFormat", {}).get("numberFormat", {}).get("type")
    if kind in ("DATE", "DATE_TIME"):
        moment = _SHEETS_EPOCH + dt.timedelta(days=number)
        moment = moment.replace(microsecond=round(moment.microsecond / 1000) * 1000 % 1_000_000)
        return moment.date() if kind == "DATE" else moment
    if kind == "TIME":
        seconds = round((number % 1) * 86400)
        return (dt.datetime.min + dt.timedelta(seconds=seconds)).time()
    if float(number).is_integer() and abs(number) < 2**53:
        return int(number)
    return number


def parse_sheets_api(payload: dict[str, Any]) -> tuple[str, list[Tab]]:
    tabs = []
    for position, sheet in enumerate(payload.get("sheets", [])):
        props = sheet.get("properties", {})
        rows: list[list[Any]] = []
        for block in sheet.get("data", []):
            for row in block.get("rowData", []):
                rows.append([cell_value(c) for c in row.get("values", [])])
        tabs.append(tab_from_rows(props.get("title", f"Sheet{position + 1}"), props.get("index", position),
                                  bool(props.get("hidden")), rows))
    return payload.get("properties", {}).get("title") or "google_sheet", tabs


def _google_message(response: httpx.Response) -> str:
    try:
        return response.json().get("error", {}).get("message", "")
    except ValueError:
        return ""


def fetch_workbook_signed_in(
    url: str, access_token: str, client: httpx.Client | None = None, max_bytes: int = 50 * 1024 * 1024
) -> Workbook:
    """Read a sheet through the Sheets API as the person whose access token this is."""
    spreadsheet_id, published = parse_sheet_url(url)
    if published:
        # Published-to-web ids are not spreadsheet ids; those sheets are public anyway.
        return fetch_workbook(url, client=client, max_bytes=max_bytes)
    own_client = client is None
    client = client or httpx.Client(timeout=90)
    try:
        with client.stream(
            "GET",
            f"{SHEETS_API}/{spreadsheet_id}",
            params={"includeGridData": "true", "fields": SHEETS_FIELDS},
            headers={"Authorization": f"Bearer {access_token}"},
        ) as response:
            if response.status_code != 200:
                response.read()
                detail = _google_message(response)
                if response.status_code == 401:
                    raise SheetError("your Google sign-in has expired; sign in again", needs_sign_in=True)
                if response.status_code == 403 and ("has not been used" in detail or "is disabled" in detail):
                    raise SheetError(f"the Google Sheets API is not enabled for eyetable's Google project: {detail}")
                if response.status_code in (403, 404):
                    raise SheetError(
                        "your Google account can't open that spreadsheet. Check the link, or ask its owner to share it with you.",
                        needs_sign_in=True,
                    )
                if response.status_code == 400 and "not supported for this document" in detail:
                    raise SheetError("that link is an Excel file stored in Drive, not a Google Sheet. Open it in Sheets and use File > Save as Google Sheets.")
                raise SheetError(f"Google answered {response.status_code}: {detail or response.reason_phrase}")
            chunks, size = [], 0
            for chunk in response.iter_bytes():
                size += len(chunk)
                if size > max_bytes:
                    raise SheetError(f"the spreadsheet is larger than {max_bytes // (1024 * 1024)} MB")
                chunks.append(chunk)
    except httpx.HTTPError as exc:
        raise SheetError(f"couldn't reach Google Sheets ({type(exc).__name__})") from None
    finally:
        if own_client:
            client.close()
    title, tabs = parse_sheets_api(json.loads(b"".join(chunks)))
    return Workbook(title, spreadsheet_id, False, sheet_link(spreadsheet_id, False), tabs)


def fetch_workbook(url: str, client: httpx.Client | None = None, max_bytes: int = 50 * 1024 * 1024) -> Workbook:
    spreadsheet_id, published = parse_sheet_url(url)
    own_client = client is None
    client = client or httpx.Client(timeout=60, follow_redirects=True)
    try:
        with client.stream("GET", export_url(spreadsheet_id, published)) as response:
            content_type = response.headers.get("content-type", "")
            if response.status_code in (401, 403, 404) or (response.is_success and _XLSX not in content_type):
                raise SheetError(
                    "this spreadsheet is private. Sign in with Google to read it with your account, "
                    "or share it as \"Anyone with the link can view\".",
                    needs_sign_in=True,
                )
            if not response.is_success:
                raise SheetError(f"Google answered {response.status_code} for that spreadsheet")
            chunks, size = [], 0
            for chunk in response.iter_bytes():
                size += len(chunk)
                if size > max_bytes:
                    raise SheetError(f"the spreadsheet is larger than {max_bytes // (1024 * 1024)} MB")
                chunks.append(chunk)
            title = title_from_disposition(response.headers.get("content-disposition")) or "google_sheet"
    except httpx.HTTPError as exc:
        raise SheetError(f"couldn't reach Google Sheets ({type(exc).__name__})") from None
    finally:
        if own_client:
            client.close()
    return Workbook(title, spreadsheet_id, published, sheet_link(spreadsheet_id, published), parse_workbook(b"".join(chunks)))
